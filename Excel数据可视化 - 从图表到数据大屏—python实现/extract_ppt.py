# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Emu

path = r'C:\Users\DENG_影\Desktop\作业\数据可视化\《Excel数据可视化 - 从图表到数据大屏》-清华-郭宏远\《Excel数据可视化——从图表到数据大屏》.pptx'
p = Presentation(path)

for i, slide in enumerate(p.slides, 1):
    print(f'===== 第 {i} 页 =====')
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text.strip()
            if txt:
                print('[文本]', txt)
        if shape.shape_type == 13:  # picture
            print('[图片]', shape.name)
        if shape.has_table:
            print('[表格]')
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                print(' | '.join(cells))
        if shape.has_chart:
            chart = shape.chart
            print('[图表]', chart.chart_type)
            try:
                cats = list(chart.plots[0].categories)
                print('  分类:', cats)
                for s in chart.series:
                    print('  系列:', s.name, list(s.values))
            except Exception as e:
                print('  读取图表数据失败:', e)
    print()
